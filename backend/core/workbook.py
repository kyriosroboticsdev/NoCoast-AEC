"""The schedules as an Excel workbook: what the office opens next to the drawings.

One sheet per schedule, using the same room numbers and door and window marks as the drawing set:

    Summary      project data, areas, occupancy, code score, cost, carbon
    Areas        GIA / NIA / circulation / efficiency / occupants by storey
    Rooms        room data: number, name, use, level, area (m² and sf), perimeter, clear height, glazing, occupants
    Doors        door schedule by mark
    Windows      window schedule by mark
    Equipment    fixtures and library assets counted by type and storey, with their IFC class
    Cost plan    UniFormat II elemental estimate
    Carbon       A1–A5 upfront carbon by element
    Code review  every clause with its measured and required values and the advice

Totals, sf conversions and extended rates are live formulas, so the schedules still add up after the
office edits them. Header rows are frozen and filterable.
"""

from __future__ import annotations

import io
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from core.draw.sheets import marks
from schemas.bim import Asset, BuildingSpec, CustomFixture, Door, Fixture, Wall, Window

SF = 10.7639
HEAD = PatternFill("solid", fgColor="1F2937")
BAND = PatternFill("solid", fgColor="F3F4F6")
TOTAL = PatternFill("solid", fgColor="E5E7EB")
STATUS = {"pass": ("PASS", "DCFCE7"), "warn": ("REVIEW", "FEF3C7"), "fail": ("FAIL", "FEE2E2"), "info": ("INFO", "E0F2FE")}
THIN = Side(style="thin", color="D1D5DB")


def _table(ws: Worksheet, title: str, cols: list[tuple[str, int, str | None]], rows: list[list], total: dict[int, str] | None = None,
           start: int = 1) -> int:
    """Title, a frozen filterable header and the rows; `total` maps a column index to "sum" (or a label).
    `cols` is (header, width, number format). Returns the row after the table."""
    ws.cell(start, 1, title).font = Font(bold=True, size=13)
    head = start + 2
    for j, (name, width, _) in enumerate(cols, 1):
        c = ws.cell(head, j, name)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = max(ws.column_dimensions[get_column_letter(j)].width or 0, width)
    for i, row in enumerate(rows):
        r = head + 1 + i
        for j, value in enumerate(row, 1):
            c = ws.cell(r, j, value)
            fmt = cols[j - 1][2]
            if fmt:
                c.number_format = fmt
            c.border = Border(bottom=THIN)
            if i % 2:
                c.fill = BAND
    last = head + len(rows)
    if start == 1:
        ws.freeze_panes = ws.cell(head + 1, 1)
        if rows:
            ws.auto_filter.ref = f"A{head}:{get_column_letter(len(cols))}{last}"
    if total and rows:
        r = last + 1
        for j in range(1, len(cols) + 1):
            c = ws.cell(r, j)
            c.fill, c.font = TOTAL, Font(bold=True)
            how = total.get(j)
            if how == "sum":
                col = get_column_letter(j)
                c.value = f"=SUM({col}{head + 1}:{col}{last})"
                c.number_format = cols[j - 1][2] or "General"
            elif how:
                c.value = how
        last = r
    return last + 2


def _level_names(spec: BuildingSpec) -> dict[str, str]:
    return {l.id: l.name for l in spec.levels}


def _summary(ws: Worksheet, spec: BuildingSpec, meta: dict, rev: dict, est: dict) -> None:
    occ, tot, cost, carbon, score = rev["occupancy"], rev["totals"], est["cost"], est["carbon"], rev["score"]
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 30, 70
    ws.cell(1, 1, spec.building.name).font = Font(bold=True, size=16)
    ws.cell(2, 1, meta.get("status", "Preliminary — not for construction")).font = Font(italic=True, color="B45309")
    rows = [
        ("Project", f"{meta.get('project_id', '')} · version {meta.get('version', 1)}"),
        ("Issued", meta.get("date", "")),
        ("Brief", meta.get("brief", "")),
        ("Code basis", rev["code"]),
        ("Occupancy", f"{occ['group']} — {occ['name']}" + (f" (with {', '.join(occ['mixed'])})" if occ.get("mixed") else "")),
        ("Design occupant load", occ["load"]),
        ("Construction", occ["construction"]),
        ("Storeys", tot["storeys"]),
        ("Gross internal area (m²)", tot["gia"]),
        ("Gross internal area (sf)", tot["gia_sf"]),
        ("Net internal area (m²)", tot["nia"]),
        ("Efficiency (NIA / GIA)", tot["efficiency"]),
        ("Code screen", f"{score['pass']} pass · {score['warn']} to review · {score['fail']} fail of {score['total']} clauses"),
        ("Construction cost (USD)", cost["total"]),
        ("Cost range (USD)", f"{cost['low']:,} – {cost['high']:,} ({cost['class']})"),
        ("Cost per m² / sf (USD)", f"{cost['per_m2']:,} / {cost['per_sf']:,}"),
        ("Upfront carbon (kgCO₂e/m²)", carbon["per_m2"]),
        ("LETI band", f"{carbon['band']} ({carbon['typology']}; 2030 target {carbon['target_2030']})"),
    ]
    fmts = {"Gross internal area (m²)": "#,##0.0", "Net internal area (m²)": "#,##0.0", "Gross internal area (sf)": "#,##0",
            "Efficiency (NIA / GIA)": "0.0%", "Construction cost (USD)": "$#,##0"}
    for i, (k, v) in enumerate(rows, 4):
        ws.cell(i, 1, k).font = Font(bold=True)
        c = ws.cell(i, 2, v)
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
        if k in fmts:
            c.number_format = fmts[k]


def _areas(ws: Worksheet, rev: dict) -> None:
    cols = [("Level", 22, None), ("GIA m²", 12, "#,##0.0"), ("NIA m²", 12, "#,##0.0"), ("Circulation m²", 14, "#,##0.0"),
            ("NIA / GIA", 11, "0.0%"), ("Rooms", 9, "0"), ("Occupants", 11, "0"), ("GIA sf", 12, "#,##0")]
    rows = [[l["name"], l["gia"], l["nia"], l["circulation"], l["efficiency"], l["rooms"], l["occupants"], None]
            for l in rev["levels"]]
    end = _table(ws, "Area summary", cols, rows, {1: "TOTAL", 2: "sum", 3: "sum", 4: "sum", 6: "sum", 7: "sum", 8: "sum"})
    first = 4
    for i in range(len(rows)):
        r = first + i
        ws.cell(r, 8, f"=B{r}*{SF}").number_format = "#,##0"
    if rows:
        t = first + len(rows)
        ws.cell(t, 5, f"=C{t}/B{t}").number_format = "0.0%"
    ws.cell(end, 1, "GIA and NIA measured to the internal face of external walls; circulation is halls, corridors, "
                    "landings and stair halls.").font = Font(italic=True, color="6B7280")


def _rooms(ws: Worksheet, rev: dict) -> None:
    cols = [("No.", 7, None), ("Room", 26, None), ("Use", 13, None), ("Level", 16, None), ("Area m²", 10, "#,##0.00"),
            ("Area sf", 10, "#,##0"), ("Perimeter m", 12, "#,##0.0"), ("Clear height mm", 14, "#,##0"),
            ("Glazing m²", 11, "#,##0.0"), ("Glazing / floor", 13, "0%"), ("Doors", 7, "0"), ("Windows", 9, "0"),
            ("Occupants", 11, "0")]
    rows = [[r["number"], r["name"], r["kind"], r["level_name"], r["area"], None, r["perimeter"], round(r["clear_height"] * 1000),
             r["glazing"], r["window_floor_ratio"], r["doors"], r["windows"], r["occupants"] or 0] for r in rev["rooms"]]
    _table(ws, "Room schedule", cols, rows, {1: "TOTAL", 5: "sum", 6: "sum", 9: "sum", 11: "sum", 12: "sum", 13: "sum"})
    for i in range(len(rows)):
        ws.cell(4 + i, 6, f"=E{4 + i}*{SF}").number_format = "#,##0"


def _openings(ws_doors: Worksheet, ws_windows: Worksheet, spec: BuildingSpec) -> None:
    tag, names = marks(spec), _level_names(spec)
    walls = {e.id: e for e in spec.elements if isinstance(e, Wall)}
    doors, windows = [], []
    for e in spec.elements:
        w = walls.get(getattr(e, "wall", ""))
        level = names.get(w.level, "") if w else ""
        where = "Exterior" if w and w.external else "Interior"
        if isinstance(e, Door):
            doors.append([tag.get(e.id, e.id), e.kind.title(), round(e.width * 1000), round(e.height * 1000), level, where,
                          e.name or e.id])
        elif isinstance(e, Window):
            windows.append([tag.get(e.id, e.id), round(e.width * 1000), round(e.height * 1000), round(e.sill_height * 1000),
                            None, level, where, e.name or e.id])
    doors.sort(key=lambda r: r[0])
    windows.sort(key=lambda r: r[0])
    _table(ws_doors, "Door schedule", [("Mark", 8, None), ("Type", 11, None), ("Width mm", 10, "#,##0"), ("Height mm", 10, "#,##0"),
                                       ("Level", 16, None), ("Location", 10, None), ("Description", 44, None)], doors,
           {1: f"{len(doors)} doors"})
    _table(ws_windows, "Window schedule", [("Mark", 8, None), ("Width mm", 10, "#,##0"), ("Height mm", 10, "#,##0"),
                                           ("Sill mm", 9, "#,##0"), ("Area m²", 9, "#,##0.00"), ("Level", 16, None),
                                           ("Location", 10, None), ("Description", 44, None)], windows,
           {1: f"{len(windows)} windows", 5: "sum"})
    for i in range(len(windows)):
        r = 4 + i
        ws_windows.cell(r, 5, f"=B{r}*C{r}/1000000").number_format = "#,##0.00"


def _equipment(ws: Worksheet, spec: BuildingSpec) -> None:
    names = _level_names(spec)
    count: Counter = Counter()
    for e in spec.elements:
        if isinstance(e, Fixture):
            count[(e.kind.replace("_", " ").capitalize(), "IfcFurnishingElement", names.get(e.level, e.level))] += 1
        elif isinstance(e, CustomFixture):
            count[(e.name or "Purpose-made piece", "IfcFurnishingElement", names.get(e.level, e.level))] += 1
        elif isinstance(e, Asset):
            count[((e.name or e.brick).split(" — ")[0], e.ifc_class, names.get(e.level, e.level))] += 1
        elif e.type in ("outlet", "light", "panel"):
            label = {"outlet": "Power outlet", "panel": "Distribution board"}.get(e.type, "Light fitting")
            count[(label, "IfcFlowTerminal" if e.type != "panel" else "IfcElectricDistributionBoard",
                   names.get(getattr(e, "level", ""), ""))] += 1
    rows = [[k[0], k[1], k[2], n] for k, n in sorted(count.items(), key=lambda kv: (kv[0][2], kv[0][0]))]
    _table(ws, "Equipment and fittings", [("Item", 34, None), ("IFC class", 26, None), ("Level", 16, None), ("Count", 8, "0")],
           rows, {1: "TOTAL", 4: "sum"})


def _cost(ws: Worksheet, est: dict) -> None:
    cost = est["cost"]
    cols = [("UniFormat", 10, None), ("Element", 44, None), ("Quantity", 11, "#,##0.0"), ("Unit", 7, None),
            ("Rate USD", 11, "$#,##0"), ("Total USD", 14, "$#,##0")]
    rows = [[f"{l['group']} {l['group_name']}", l["element"], l["quantity"], l["unit"], l["rate"], None] for l in cost["lines"]]
    rows += [["", "General conditions, OH&P", None, "", None, cost["general_conditions"]],
             ["", "Design contingency", None, "", None, cost["contingency"]]]
    end = _table(ws, f"Cost plan — {cost['class']}", cols, rows, {2: "TOTAL", 6: "sum"})
    for i in range(len(cost["lines"])):
        r = 4 + i
        ws.cell(r, 6, f"=C{r}*E{r}").number_format = "$#,##0"
    ws.cell(end, 1, cost["basis"]).font = Font(italic=True, color="6B7280")


def _carbon(ws: Worksheet, est: dict) -> None:
    carbon = est["carbon"]
    rows = [[r["element"], r["kg"], None] for r in carbon["rows"]]
    end = _table(ws, "Upfront embodied carbon, A1–A5", [("Element", 30, None), ("kgCO₂e", 14, "#,##0"), ("Share", 9, "0%")],
                 rows, {1: "TOTAL", 2: "sum"})
    t = 4 + len(rows)
    for i in range(len(rows)):
        ws.cell(4 + i, 3, f"=B{4 + i}/B${t}").number_format = "0%"
    ws.cell(end, 1, f"{carbon['per_m2']} kgCO₂e/m², LETI band {carbon['band']} — {carbon['basis']}").font = Font(italic=True, color="6B7280")
    opts = [[o["move"], -o["saving_kg"], o["saving_pct"], o["per_m2"]] for o in carbon["options"]]
    if opts:
        _table(ws, "Design moves that cut carbon", [("Move", 30, None), ("kgCO₂e", 14, "#,##0"), ("Share", 9, "0%"),
                                                     ("Then kgCO₂e/m²", 16, "0")], opts, start=end + 2)


def _review(ws: Worksheet, rev: dict) -> None:
    cols = [("Status", 9, None), ("Clause", 16, None), ("Check", 34, None), ("Measured", 40, None), ("Required", 40, None),
            ("Advice", 60, None)]
    rows = [[STATUS[c["status"]][0], c["reference"], c["title"], c["value"], c["target"], c["advice"] or ""] for c in rev["checks"]]
    _table(ws, f"Code review — {rev['code']}", cols, rows)
    for i, c in enumerate(rev["checks"]):
        ws.cell(4 + i, 1).fill = PatternFill("solid", fgColor=STATUS[c["status"]][1])
        for j in (4, 5, 6):
            ws.cell(4 + i, j).alignment = Alignment(wrap_text=True, vertical="top")


def schedules_xlsx(spec: BuildingSpec, rev: dict, est: dict, meta: dict) -> bytes:
    wb = Workbook()
    sheets = ["Summary", "Areas", "Rooms", "Doors", "Windows", "Equipment", "Cost plan", "Carbon", "Code review"]
    ws = {name: (wb.active if i == 0 else wb.create_sheet()) for i, name in enumerate(sheets)}
    for name, sheet in ws.items():
        sheet.title = name
    _summary(ws["Summary"], spec, meta, rev, est)
    _areas(ws["Areas"], rev)
    _rooms(ws["Rooms"], rev)
    _openings(ws["Doors"], ws["Windows"], spec)
    _equipment(ws["Equipment"], spec)
    _cost(ws["Cost plan"], est)
    _carbon(ws["Carbon"], est)
    _review(ws["Code review"], rev)
    wb.properties.title = f"{spec.building.name} — schedules"
    wb.properties.creator = "NoCoast AEC"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
