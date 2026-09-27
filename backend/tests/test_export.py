"""Exporting a version: the IFC on its own, the individual artefacts, and the whole bundle."""

import csv
import io
import json
import zipfile

import ezdxf
import ifcopenshell
from ezdxf import recover
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from main import app
from tests.sse import done

client = TestClient(app)


def build() -> tuple[str, dict]:
    pid = client.post("/projects", json={"name": "export"}).json()["id"]
    version = done(client.post(f"/projects/{pid}/prompt",
                               json={"prompt": "Two storey house with a kitchen, living room and two bedrooms"}).text)
    return pid, version


def test_version_carries_its_export_url():
    pid, version = build()
    assert version["export_url"] == f"/projects/{pid}/versions/1/export"
    assert client.get(f"/projects/{pid}").json()["head"]["export_url"] == version["export_url"]


def test_ifc_download_is_named_after_the_version():
    pid, _ = build()
    r = client.get(f"/projects/{pid}/versions/1/ifc")
    assert r.status_code == 200
    assert f'filename="{pid}-v1.ifc"' in r.headers["content-disposition"]
    assert ifcopenshell.file.from_string(r.text).by_type("IfcProject")


def test_single_artifacts():
    pid, _ = build()

    spec = client.get(f"/projects/{pid}/versions/1/export", params={"format": "spec"})
    assert spec.status_code == 200 and json.loads(spec.content)["levels"]

    design = client.get(f"/projects/{pid}/versions/1/export", params={"format": "design"})
    assert json.loads(design.content)["rooms"]

    context = client.get(f"/projects/{pid}/versions/1/export", params={"format": "context"})
    assert "room id=kitchen" in context.text

    summary = client.get(f"/projects/{pid}/versions/1/export", params={"format": "summary"})
    assert summary.text.startswith("# ") and "## Design approach" in summary.text and "## Requirements" in summary.text

    schedule = client.get(f"/projects/{pid}/versions/1/export", params={"format": "schedule"})
    rows = list(csv.DictReader(io.StringIO(schedule.text)))
    assert {r["category"] for r in rows} >= {"room", "wall", "door", "window"}
    kitchen = next(r for r in rows if r["category"] == "room" and r["name"] == "Kitchen")
    assert float(kitchen["quantity"]) > 5 and kitchen["unit"] == "m2"

    assert client.get(f"/projects/{pid}/versions/1/export", params={"format": "nope"}).status_code == 400


def test_bundle_holds_everything():
    pid, _ = build()
    r = client.get(f"/projects/{pid}/versions/1/export")
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    assert f'filename="{pid}-v1.zip"' in r.headers["content-disposition"]

    with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
        names = set(zf.namelist())
        drawings = {n for n in names if n.startswith("drawings/")}
        assert drawings >= {"drawings/G-001.svg", "drawings/A-101.svg", "drawings/A-201.svg", "drawings/A-301.svg"}
        cad = {n for n in names if n.startswith("cad/")}
        levels = json.loads(zf.read(f"{pid}-v1.spec.json"))["levels"]
        assert len(cad) == len(levels) == 2 and all(n.endswith(".dxf") for n in cad)
        assert names - drawings - cad == {"README.md", "guids.json"} | {
            f"{pid}-v1.{ext}" for ext in ("ifc", "drawings.pdf", "plans.dxf", "schedules.xlsx", "review.md",
                                          "issues.bcfzip", "estimate.csv", "summary.md", "spec.json", "design.json",
                                          "context.txt", "checks.json", "schedule.csv")}
        assert zf.read(f"{pid}-v1.drawings.pdf").startswith(b"%PDF-")
        assert zf.read("README.md").decode().startswith("# NoCoast export")
        model = ifcopenshell.file.from_string(zf.read(f"{pid}-v1.ifc").decode())
        assert model.by_type("IfcWall") and model.by_type("IfcSpace")
        assert json.loads(zf.read("guids.json"))


def test_plans_as_dxf_on_ncs_layers():
    pid, _ = build()
    r = client.get(f"/projects/{pid}/versions/1/export", params={"format": "dxf"})
    assert r.status_code == 200 and f'filename="{pid}-v1.plans.dxf"' in r.headers["content-disposition"]
    doc, auditor = recover.read(io.BytesIO(r.content))
    assert not auditor.has_errors
    assert doc.units == ezdxf.units.M
    msp = doc.modelspace()
    layers = {e.dxf.layer for e in msp}
    assert layers >= {"A-WALL", "A-WALL-PATT", "A-DOOR", "A-DOOR-IDEN", "A-GLAZ", "A-GLAZ-IDEN", "A-AREA", "A-AREA-IDEN",
                      "A-FLOR-STRS", "S-GRID", "A-ANNO-DIMS", "A-ANNO-TTLB"}
    rooms = [e for e in msp.query("LWPOLYLINE[layer=='A-AREA']")]
    spec = json.loads(client.get(f"/projects/{pid}/versions/1/export", params={"format": "spec"}).content)
    assert len(rooms) == sum(1 for e in spec["elements"] if e["type"] == "space")
    tags = {e.dxf.text for e in msp.query("TEXT[layer=='A-AREA-IDEN']")}
    assert "KITCHEN" in tags and any(t.endswith("m²") for t in tags)
    titles = {e.dxf.text for e in msp.query("TEXT[layer=='A-ANNO-TTLB']")}
    assert {f"{l['name'].upper()} PLAN" for l in spec["levels"]} <= titles
    assert msp.query("DIMENSION") and msp.query("ARC[layer=='A-DOOR']")


def test_schedules_workbook():
    pid, _ = build()
    r = client.get(f"/projects/{pid}/versions/1/export", params={"format": "xlsx"})
    assert r.status_code == 200 and f'filename="{pid}-v1.schedules.xlsx"' in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Areas", "Rooms", "Doors", "Windows", "Equipment", "Cost plan", "Carbon", "Code review"]
    analysis = client.get(f"/projects/{pid}/versions/1/analysis").json()
    rooms = wb["Rooms"]
    assert [c.value for c in rooms[3]][:5] == ["No.", "Room", "Use", "Level", "Area m²"]
    numbers = [rooms.cell(r, 1).value for r in range(4, 4 + len(analysis["review"]["rooms"]))]
    assert numbers == [room["number"] for room in analysis["review"]["rooms"]]
    total = next(r for r in range(4, rooms.max_row + 1) if rooms.cell(r, 1).value == "TOTAL")
    assert rooms.cell(total, 5).value == f"=SUM(E4:E{total - 1})" and rooms.cell(4, 6).value.startswith("=E4*")
    doors = wb["Doors"]
    marks = [doors.cell(r, 1).value for r in range(4, doors.max_row) if doors.cell(r, 1).value]
    assert marks[0].startswith("D1") and len([m for m in marks if m.startswith("D")]) >= 3
    cost = wb["Cost plan"]
    assert any(cost.cell(r, 2).value == "TOTAL" for r in range(1, cost.max_row + 1))
    statuses = {wb["Code review"].cell(r, 1).value for r in range(4, 4 + len(analysis["review"]["checks"]))}
    assert statuses <= {"PASS", "REVIEW", "FAIL", "INFO"} and statuses


def test_project_export_is_the_head_version():
    pid, _ = build()
    done(client.post(f"/projects/{pid}/prompt", json={"prompt": "add a bedroom"}).text)
    r = client.get(f"/projects/{pid}/export")
    assert f'filename="{pid}-v2.zip"' in r.headers["content-disposition"]
    with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
        assert f"{pid}-v2.ifc" in zf.namelist()
    assert client.get("/projects/nope/export").status_code == 404


def test_exported_ifc_can_be_imported_back():
    pid, _ = build()
    with zipfile.ZipFile(io.BytesIO(client.get(f"/projects/{pid}/versions/1/export").content)) as zf:
        ifc = zf.read(f"{pid}-v1.ifc")
    other = client.post("/projects", json={"name": "round trip"}).json()["id"]
    v = done(client.post(f"/projects/{other}/import", files={"file": ("house.ifc", ifc, "application/x-step")}).text)
    assert v["mode"] == "import" and v["summary"]["elements"] > 10
